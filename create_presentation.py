"""
ScamShield AI — Premium ForgeHacks 2026 Presentation Deck Generator
Creates a professional 16:9 widescreen presentation in PPTX format.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette (Dark Cybersecurity SaaS Theme)
BG_DARK = RGBColor(7, 11, 20)          # #070B14 Near-black navy
CARD_BG = RGBColor(13, 20, 36)         # #0D1424 Dark navy card
CARD_BG_ALT = RGBColor(16, 26, 48)     # #101A30 Slightly lighter card
CARD_BORDER = RGBColor(30, 58, 95)     # #1E3A5F Subtle blue border
BORDER_CYAN = RGBColor(0, 242, 254)    # #00F2FE Electric cyan accent
CYAN = RGBColor(0, 242, 254)           # #00F2FE Primary accent
SKY = RGBColor(56, 189, 248)           # #38BDF8 Secondary cyan
TEXT_WHITE = RGBColor(255, 255, 255)   # #FFFFFF Clean white
TEXT_SILVER = RGBColor(203, 213, 225)  # #CBD5E1 High contrast silver
TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B Secondary muted
RED = RGBColor(239, 68, 68)            # #EF4444 High risk red
RED_BG = RGBColor(38, 16, 22)          # #261016 Dark red tint
AMBER = RGBColor(245, 158, 11)         # #F59E0B Warning amber
GREEN = RGBColor(16, 185, 129)         # #10B981 Safe green
GREEN_BG = RGBColor(10, 35, 26)        # #0A231A Dark green tint

FONT_MAIN = "Helvetica Neue"
FONT_MONO = "Courier New"


def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_header(slide, slide_num, total_slides, category):
    # Brand tag on top left
    tb_brand = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(5.0), Inches(0.35))
    tf = tb_brand.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "SCAMSHIELD AI"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN

    run = p.add_run()
    run.text = f"  •  {category.upper()}"
    run.font.name = FONT_MAIN
    run.font.size = Pt(10)
    run.font.bold = False
    run.font.color.rgb = TEXT_MUTED

    # Slide number on top right
    tb_num = slide.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.0), Inches(0.35))
    tf_num = tb_num.text_frame
    tf_num.word_wrap = False
    tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
    p_num = tf_num.paragraphs[0]
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.text = f"{slide_num:02d} / {total_slides:02d}"
    p_num.font.name = FONT_MONO
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = SKY


def add_footer(slide):
    # Professional subtle footer line
    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = tb_foot.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "FORGEHACKS 2026  •  AI + CYBERSECURITY TRACK  •  51/51 TESTS PASSED"
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED


def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape


def build_presentation(output_path="ScamShield_AI_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    TOTAL_SLIDES = 6

    # =========================================================================
    # SLIDE 1: HERO
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, 1, TOTAL_SLIDES, "Opening Showcase")
    add_footer(s1)

    # Hero Shield Visual / Glowing Container
    add_card(s1, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.6), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1.5)

    # Main Headline
    tb_hero = s1.shapes.add_textbox(Inches(1.3), Inches(1.6), Inches(10.7), Inches(2.2))
    tf_hero = tb_hero.text_frame
    tf_hero.word_wrap = True
    tf_hero.margin_left = tf_hero.margin_top = tf_hero.margin_right = tf_hero.margin_bottom = 0

    p_tag = tf_hero.paragraphs[0]
    p_tag.text = "NEXT-GENERATION CONSUMER THREAT INTERCEPTION"
    p_tag.font.name = FONT_MONO
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = SKY

    p_title = tf_hero.add_paragraph()
    p_title.text = "SCAMSHIELD AI"
    p_title.font.name = FONT_MAIN
    p_title.font.size = Pt(48)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    p_sub = tf_hero.add_paragraph()
    p_sub.text = "AI-Powered Scam Detection & Verification"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = CYAN

    # Supporting line
    tb_sup = s1.shapes.add_textbox(Inches(1.3), Inches(3.9), Inches(10.7), Inches(0.8))
    tf_sup = tb_sup.text_frame
    tf_sup.word_wrap = True
    tf_sup.margin_left = tf_sup.margin_top = tf_sup.margin_right = tf_sup.margin_bottom = 0
    p_sup = tf_sup.paragraphs[0]
    p_sup.text = '“Identify risk. Understand why. Act safely.”'
    p_sup.font.name = FONT_MAIN
    p_sup.font.size = Pt(17)
    p_sup.font.italic = True
    p_sup.font.color.rgb = TEXT_SILVER

    p_sup2 = tf_sup.add_paragraph()
    p_sup2.text = "A zero-harm verification engine that inspects suspicious messages, extracted URLs, and screenshots before users click."
    p_sup2.font.name = FONT_MAIN
    p_sup2.font.size = Pt(13)
    p_sup2.font.color.rgb = TEXT_MUTED

    # Badge Row (4 Pill Badges)
    badges = [
        ("AI + CYBERSECURITY", SKY),
        ("FORGEHACKS 2026", CYAN),
        ("GEMINI 2.5 FLASH", TEXT_SILVER),
        ("51/51 TESTS PASSED", GREEN),
    ]
    b_width = Inches(2.55)
    b_gap = Inches(0.2)
    b_left = Inches(1.3)
    for i, (b_text, b_color) in enumerate(badges):
        bx = b_left + i * (b_width + b_gap)
        add_card(s1, bx, Inches(5.3), b_width, Inches(0.65), bg_color=CARD_BG_ALT, border_color=CARD_BORDER, border_width=1)
        tb_b = s1.shapes.add_textbox(bx, Inches(5.3), b_width, Inches(0.65))
        tf_b = tb_b.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_b.word_wrap = False
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = b_text
        p_b.font.name = FONT_MONO
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = b_color

    # =========================================================================
    # SLIDE 2: THE PROBLEM
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, 2, TOTAL_SLIDES, "The Threat Landscape")
    add_footer(s2)

    # Slide Title
    tb_t2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf_t2 = tb_t2.text_frame
    tf_t2.word_wrap = True
    p2 = tf_t2.paragraphs[0]
    p2.text = "The Social Engineering Epidemic"
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2_sub = tf_t2.add_paragraph()
    p2_sub.text = "Attackers exploit psychological urgency across channels where legacy filters are blind."
    p2_sub.font.name = FONT_MAIN
    p2_sub.font.size = Pt(14)
    p2_sub.font.color.rgb = TEXT_MUTED

    # 4 Threat Cards
    threats = [
        ("PHISHING", "🚨 URGENCY & PANIC", "Fake KYC alerts, account freeze deadlines, and false security warnings engineered to trigger immediate action.", RED),
        ("IMPERSONATION", "🎭 FALSE AUTHORITY", "Spoofed bank notices, courier package holds, and government agency identities designed to steal credentials.", SKY),
        ("MALICIOUS LINKS", "🎣 WEAPONIZED URLS", "Raw IP hosts, misleading subdomains, userinfo @ spoofing, and high-abuse TLDs (.xyz, .top) hiding destinations.", AMBER),
        ("SCAM SCREENSHOTS", "📸 THE BLIND SPOT", "Suspicious messages forwarded as screenshot images across chat apps where standard text filters cannot read.", CYAN),
    ]

    card_w = Inches(2.75)
    gap_w = Inches(0.24)
    start_x = Inches(0.8)

    for i, (title, tag, desc, color) in enumerate(threats):
        cx = start_x + i * (card_w + gap_w)
        add_card(s2, cx, Inches(1.85), card_w, Inches(3.2), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)

        tb_card = s2.shapes.add_textbox(cx + Inches(0.25), Inches(2.05), card_w - Inches(0.5), Inches(2.8))
        tfc = tb_card.text_frame
        tfc.word_wrap = True
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = 0

        p_tag = tfc.paragraphs[0]
        p_tag.text = tag
        p_tag.font.name = FONT_MONO
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = color

        p_head = tfc.add_paragraph()
        p_head.text = title
        p_head.font.name = FONT_MAIN
        p_head.font.size = Pt(18)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_WHITE
        p_head.space_before = Pt(8)
        p_head.space_after = Pt(10)

        p_desc = tfc.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_SILVER

    # Emotional Problem Statement Card (Bottom callout)
    add_card(s2, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.4), bg_color=CARD_BG_ALT, border_color=CYAN, border_width=1.5)
    tb_q = s2.shapes.add_textbox(Inches(1.2), Inches(5.45), Inches(10.933), Inches(1.2))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    tf_q.margin_left = tf_q.margin_top = tf_q.margin_right = tf_q.margin_bottom = 0

    pq1 = tf_q.paragraphs[0]
    pq1.text = '“Is this safe to interact with?”'
    pq1.font.name = FONT_MAIN
    pq1.font.size = Pt(22)
    pq1.font.bold = True
    pq1.font.color.rgb = CYAN

    pq2 = tf_q.add_paragraph()
    pq2.text = "Everyday users have no safe sandbox to verify suspicious messages. Testing a link by clicking it is dangerous, and existing filters provide zero explanation."
    pq2.font.name = FONT_MAIN
    pq2.font.size = Pt(13)
    pq2.font.color.rgb = TEXT_SILVER
    pq2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: THE SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, 3, TOTAL_SLIDES, "The Solution")
    add_footer(s3)

    # Title
    tb_t3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf_t3 = tb_t3.text_frame
    tf_t3.word_wrap = True
    p3 = tf_t3.paragraphs[0]
    p3.text = "End-to-End Threat Interception Pipeline"
    p3.font.name = FONT_MAIN
    p3.font.size = Pt(28)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_WHITE
    p3_sub = tf_t3.add_paragraph()
    p3_sub.text = "Multimodal ingestion couples neural contextual evaluation with deterministic zero-network URL inspection."
    p3_sub.font.name = FONT_MAIN
    p3_sub.font.size = Pt(14)
    p3_sub.font.color.rgb = TEXT_MUTED

    # 6-Node Horizontal Pipeline
    pipe_nodes = [
        ("01. INGESTION", "Message or\nScreenshot", SKY),
        ("02. OCR EXTRACT", "Gemini Vision\nIn-Memory RAM", CYAN),
        ("03. SCAMSHIELD", "Detection\nOrchestrator", TEXT_WHITE),
        ("04. URL INSPECT", "Static Lexical\nZero-Network", AMBER),
        ("05. ASSESSMENT", "Calibrated\nRisk Score", RED),
        ("06. ACTIONS", "Prescriptive\nSafe Guidance", GREEN),
    ]

    node_w = Inches(1.8)
    node_h = Inches(1.6)
    node_gap = Inches(0.186)
    start_nx = Inches(0.8)

    for i, (step, label, col) in enumerate(pipe_nodes):
        nx = start_nx + i * (node_w + node_gap)
        add_card(s3, nx, Inches(1.9), node_w, node_h, bg_color=CARD_BG, border_color=col, border_width=1.2)

        tb_n = s3.shapes.add_textbox(nx + Inches(0.1), Inches(2.05), node_w - Inches(0.2), node_h - Inches(0.3))
        tfn = tb_n.text_frame
        tfn.word_wrap = True
        tfn.margin_left = tfn.margin_top = tfn.margin_right = tfn.margin_bottom = 0

        p_step = tfn.paragraphs[0]
        p_step.alignment = PP_ALIGN.CENTER
        p_step.text = step
        p_step.font.name = FONT_MONO
        p_step.font.size = Pt(9.5)
        p_step.font.bold = True
        p_step.font.color.rgb = col

        p_lbl = tfn.add_paragraph()
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.text = label
        p_lbl.font.name = FONT_MAIN
        p_lbl.font.size = Pt(12)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = TEXT_WHITE
        p_lbl.space_before = Pt(8)

    # Security Callout Banner
    add_card(s3, Inches(0.8), Inches(3.75), Inches(11.733), Inches(0.8), bg_color=CARD_BG_ALT, border_color=CYAN, border_width=1)
    tb_sec = s3.shapes.add_textbox(Inches(1.2), Inches(3.85), Inches(10.933), Inches(0.6))
    tf_sec = tb_sec.text_frame
    tf_sec.word_wrap = True
    p_sec = tf_sec.paragraphs[0]
    p_sec.text = "🔒 ZERO-HARM SECURITY PRINCIPLE:  “Suspicious URLs are analyzed without visiting them.”"
    p_sec.font.name = FONT_MONO
    p_sec.font.size = Pt(12.5)
    p_sec.font.bold = True
    p_sec.font.color.rgb = CYAN
    p_sec_sub = tf_sec.add_paragraph()
    p_sec_sub.text = "No outbound network requests • Defanged previews (hxxp:// and [.]) • Zero threat detonation"
    p_sec_sub.font.name = FONT_MAIN
    p_sec_sub.font.size = Pt(11.5)
    p_sec_sub.font.color.rgb = TEXT_SILVER

    # 4 Output Dossier Feature Cards
    outputs = [
        ("RISK SCORE", "0 to 100 Threat Index with clear High / Suspicious / Low visual tiering."),
        ("THREAT INDICATORS", "Specific urgency triggers, lexical anomalies, and authority spoofing cues."),
        ("EXPLANATION", "Plain-language root cause answering “Why is this message suspicious?”"),
        ("SAFE ACTIONS", "Concrete, prescriptive next steps—what to avoid and how to verify."),
    ]
    out_w = Inches(2.75)
    for i, (otitle, odesc) in enumerate(outputs):
        ox = start_nx + i * (out_w + gap_w)
        add_card(s3, ox, Inches(4.75), out_w, Inches(2.0), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)

        tb_o = s3.shapes.add_textbox(ox + Inches(0.2), Inches(4.9), out_w - Inches(0.4), Inches(1.7))
        tfo = tb_o.text_frame
        tfo.word_wrap = True
        tfo.margin_left = tfo.margin_top = tfo.margin_right = tfo.margin_bottom = 0

        po1 = tfo.paragraphs[0]
        po1.text = otitle
        po1.font.name = FONT_MONO
        po1.font.size = Pt(11)
        po1.font.bold = True
        po1.font.color.rgb = SKY

        po2 = tfo.add_paragraph()
        po2.text = odesc
        po2.font.name = FONT_MAIN
        po2.font.size = Pt(12)
        po2.font.color.rgb = TEXT_SILVER
        po2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 4: HOW IT WORKS (ARCHITECTURE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, 4, TOTAL_SLIDES, "How It Works")
    add_footer(s4)

    # Title
    tb_t4 = s4.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf_t4 = tb_t4.text_frame
    tf_t4.word_wrap = True
    p4 = tf_t4.paragraphs[0]
    p4.text = "Layered Architecture & Defense-in-Depth"
    p4.font.name = FONT_MAIN
    p4.font.size = Pt(28)
    p4.font.bold = True
    p4.font.color.rgb = TEXT_WHITE
    p4_sub = tf_t4.add_paragraph()
    p4_sub.text = "Distinguishing generative contextual intelligence from deterministic security safeguards."
    p4_sub.font.name = FONT_MAIN
    p4_sub.font.size = Pt(14)
    p4_sub.font.color.rgb = TEXT_MUTED

    # Left: 5 Horizontal Architecture Layers
    layers = [
        ("INPUT LAYER", "Text Message (SMS, Chat, Email)  +  Screenshot Ingestion (Dropzone / Cmd+V Paste)", SKY),
        ("INTELLIGENCE LAYER", "Gemini 2.5 Flash (Contextual Deception)  +  Gemini Vision (In-Memory OCR)", CYAN),
        ("SECURITY ANALYSIS LAYER", "RuleBasedAnalyzer (Keyword Heuristics)  +  Static URL Analyzer (Zero-Network)", AMBER),
        ("RESILIENCE LAYER", "Dual-Tier Fallback (Cloud AI Failure Failover)  +  Local Tesseract OCR Fallback", GREEN),
        ("OUTPUT LAYER", "Calibrated Risk Score  •  Defanged URLs  •  Recommendations  •  sessionStorage History", TEXT_SILVER),
    ]

    ly_w = Inches(7.8)
    ly_h = Inches(0.85)
    ly_gap = Inches(0.16)
    start_ly = Inches(1.85)

    for i, (ltitle, ldesc, lcol) in enumerate(layers):
        ly = start_ly + i * (ly_h + ly_gap)
        add_card(s4, Inches(0.8), ly, ly_w, ly_h, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)

        tb_ly = s4.shapes.add_textbox(Inches(1.05), ly + Inches(0.12), ly_w - Inches(0.4), ly_h - Inches(0.24))
        tfly = tb_ly.text_frame
        tfly.word_wrap = True
        tfly.margin_left = tfly.margin_top = tfly.margin_right = tfly.margin_bottom = 0

        ply1 = tfly.paragraphs[0]
        ply1.text = ltitle
        ply1.font.name = FONT_MONO
        ply1.font.size = Pt(10)
        ply1.font.bold = True
        ply1.font.color.rgb = lcol

        ply2 = tfly.add_paragraph()
        ply2.text = ldesc
        ply2.font.name = FONT_MAIN
        ply2.font.size = Pt(11.5)
        ply2.font.color.rgb = TEXT_WHITE
        ply2.space_before = Pt(3)

    # Right: Stack & Architectural Distinction Card
    rx = Inches(8.85)
    rw = Inches(3.68)
    add_card(s4, rx, Inches(1.85), rw, Inches(4.9), bg_color=CARD_BG_ALT, border_color=CARD_BORDER, border_width=1.2)

    tb_r = s4.shapes.add_textbox(rx + Inches(0.25), Inches(2.05), rw - Inches(0.5), Inches(4.5))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    tfr.margin_left = tfr.margin_top = tfr.margin_right = tfr.margin_bottom = 0

    pr1 = tfr.paragraphs[0]
    pr1.text = "CORE TECHNOLOGY STACK"
    pr1.font.name = FONT_MONO
    pr1.font.size = Pt(11)
    pr1.font.bold = True
    pr1.font.color.rgb = SKY

    stack_items = [
        ("Backend Framework", "Python 3.13 / Flask"),
        ("Official AI SDK", "google-genai (v2.28)"),
        ("Vision & Image Engine", "Pillow (In-Memory Validation)"),
        ("Static Security Engine", "ipaddress, urllib.parse"),
        ("Client Storage", "Browser sessionStorage (Max 20)"),
        ("Automated Tests", "pytest (51 Tests Passing)"),
    ]

    for label, val in stack_items:
        p_item = tfr.add_paragraph()
        p_item.text = f"• {label}: "
        p_item.font.name = FONT_MAIN
        p_item.font.size = Pt(11)
        p_item.font.bold = False
        p_item.font.color.rgb = TEXT_MUTED
        p_item.space_before = Pt(6)

        run = p_item.add_run()
        run.text = val
        run.font.bold = True
        run.font.color.rgb = TEXT_WHITE

    p_div = tfr.add_paragraph()
    p_div.text = "ENGINEERING DISTINCTIONS"
    p_div.font.name = FONT_MONO
    p_div.font.size = Pt(10.5)
    p_div.font.bold = True
    p_div.font.color.rgb = CYAN
    p_div.space_before = Pt(14)

    p_d1 = tfr.add_paragraph()
    p_d1.text = "• AI Context: Detects persuasive psychology & urgency."
    p_d1.font.name = FONT_MAIN
    p_d1.font.size = Pt(10.5)
    p_d1.font.color.rgb = TEXT_SILVER
    p_d1.space_before = Pt(4)

    p_d2 = tfr.add_paragraph()
    p_d2.text = "• Deterministic Rules: Enforces risk floors (Raw IP ≥ 85)."
    p_d2.font.name = FONT_MAIN
    p_d2.font.size = Pt(10.5)
    p_d2.font.color.rgb = TEXT_SILVER
    p_d2.space_before = Pt(3)

    p_d3 = tfr.add_paragraph()
    p_d3.text = "• Zero Storage: Screenshots never touch disk or database."
    p_d3.font.name = FONT_MAIN
    p_d3.font.size = Pt(10.5)
    p_d3.font.color.rgb = TEXT_SILVER
    p_d3.space_before = Pt(3)

    # =========================================================================
    # SLIDE 5: LIVE DEMO + DIFFERENTIATION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, 5, TOTAL_SLIDES, "Live Demo & Differentiation")
    add_footer(s5)

    # Title
    tb_t5 = s5.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf_t5 = tb_t5.text_frame
    tf_t5.word_wrap = True
    p5 = tf_t5.paragraphs[0]
    p5.text = "Product Showcase & Differentiation"
    p5.font.name = FONT_MAIN
    p5.font.size = Pt(28)
    p5.font.bold = True
    p5.font.color.rgb = TEXT_WHITE
    p5_sub = tf_t5.add_paragraph()
    p5_sub.text = "ScamShield identifies risk indicators and provides an actionable risk assessment."
    p5_sub.font.name = FONT_MAIN
    p5_sub.font.size = Pt(14)
    p5_sub.font.color.rgb = TEXT_MUTED

    # Left: 5-Step Demo Steps Bar + Result Mock Card
    left_w = Inches(6.8)

    # 5-Step Steps Bar
    add_card(s5, Inches(0.8), Inches(1.8), left_w, Inches(0.7), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)
    tb_steps = s5.shapes.add_textbox(Inches(0.9), Inches(1.85), left_w - Inches(0.2), Inches(0.6))
    tf_st = tb_steps.text_frame
    tf_st.word_wrap = True
    p_st = tf_st.paragraphs[0]
    p_st.alignment = PP_ALIGN.CENTER
    p_st.text = "1. Paste SMS  →  2. Analyze  →  3. Upload Screenshot  →  4. Inspect URL  →  5. Get Dossier"
    p_st.font.name = FONT_MONO
    p_st.font.size = Pt(10.5)
    p_st.font.bold = True
    p_st.font.color.rgb = SKY

    # Mock Product Result Card
    add_card(s5, Inches(0.8), Inches(2.65), left_w, Inches(4.1), bg_color=CARD_BG_ALT, border_color=RED, border_width=1.5)
    tb_mock = s5.shapes.add_textbox(Inches(1.05), Inches(2.8), left_w - Inches(0.5), Inches(3.8))
    tf_m = tb_mock.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0

    pm1 = tf_m.paragraphs[0]
    pm1.text = "DEMO SCAN RESULT  •  [ Illustrative Demo Result ]"
    pm1.font.name = FONT_MONO
    pm1.font.size = Pt(9.5)
    pm1.font.color.rgb = TEXT_MUTED

    pm2 = tf_m.add_paragraph()
    pm2.text = "HIGH RISK  —  88 / 100"
    pm2.font.name = FONT_MAIN
    pm2.font.size = Pt(22)
    pm2.font.bold = True
    pm2.font.color.rgb = RED
    pm2.space_before = Pt(4)

    pm3 = tf_m.add_paragraph()
    pm3.text = "DEFANGED TARGET:  hxxp://secure-bank-login[.]xyz"
    pm3.font.name = FONT_MONO
    pm3.font.size = Pt(11)
    pm3.font.bold = True
    pm3.font.color.rgb = CYAN
    pm3.space_before = Pt(6)

    pm4 = tf_m.add_paragraph()
    pm4.text = "DETECTED THREAT INDICATORS:"
    pm4.font.name = FONT_MONO
    pm4.font.size = Pt(10)
    pm4.font.bold = True
    pm4.font.color.rgb = TEXT_SILVER
    pm4.space_before = Pt(8)

    indicators = [
        "⚠️ Urgent KYC account suspension threat",
        "⚠️ High-abuse TLD (.xyz) & Unencrypted HTTP scheme",
        "⚠️ Suspicious credential harvesting path (/login)",
        "⚠️ Impersonation of trusted banking entity",
    ]
    for ind in indicators:
        p_ind = tf_m.add_paragraph()
        p_ind.text = f"  {ind}"
        p_ind.font.name = FONT_MAIN
        p_ind.font.size = Pt(10.5)
        p_ind.font.color.rgb = TEXT_WHITE
        p_ind.space_before = Pt(2)

    pm5 = tf_m.add_paragraph()
    pm5.text = "RECOMMENDED ACTION: Do not click link. Contact bank directly via verified phone."
    pm5.font.name = FONT_MAIN
    pm5.font.size = Pt(11)
    pm5.font.bold = True
    pm5.font.color.rgb = GREEN
    pm5.space_before = Pt(10)

    # Right: Why ScamShield? (Key Differentiators)
    rw5 = Inches(4.7)
    add_card(s5, Inches(7.833), Inches(1.8), rw5, Inches(4.95), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1.2)

    tb_why = s5.shapes.add_textbox(Inches(8.1), Inches(2.0), rw5 - Inches(0.5), Inches(4.6))
    tf_w = tb_why.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_top = tf_w.margin_right = tf_w.margin_bottom = 0

    pw1 = tf_w.paragraphs[0]
    pw1.text = "WHY SCAMSHIELD AI?"
    pw1.font.name = FONT_MONO
    pw1.font.size = Pt(11)
    pw1.font.bold = True
    pw1.font.color.rgb = CYAN

    diffs = [
        ("Multimodal Analysis", "Evaluates raw text, image file uploads, drag & drop, and direct Cmd+V clipboard paste."),
        ("AI + Deterministic Fallback", "Seamless continuity via RuleBasedAnalyzer when API is unavailable or rate limited."),
        ("Zero-Network URL Inspection", "Lexical analysis without visiting submitted URLs; never establishes external connections."),
        ("Strict URL Defanging", "Automatically renders links harmless (hxxp:// and [.]) to eliminate accidental clicks."),
        ("Privacy-Conscious History", "Zero permanent disk storage; scans stored ephemerally in browser sessionStorage."),
        ("51/51 Tests Passing", "Production-grade automated pytest coverage across OCR, analyzer, and URL modules."),
    ]

    for title, desc in diffs:
        pd = tf_w.add_paragraph()
        pd.text = f"• {title}: "
        pd.font.name = FONT_MAIN
        pd.font.size = Pt(11)
        pd.font.bold = True
        pd.font.color.rgb = SKY
        pd.space_before = Pt(6)

        run = pd.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_SILVER

    # =========================================================================
    # SLIDE 6: IMPACT + ROADMAP
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, 6, TOTAL_SLIDES, "Impact & Roadmap")
    add_footer(s6)

    # Title
    tb_t6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf_t6 = tb_t6.text_frame
    tf_t6.word_wrap = True
    p6 = tf_t6.paragraphs[0]
    p6.text = "Impact, Honest Limits & Future Scope"
    p6.font.name = FONT_MAIN
    p6.font.size = Pt(28)
    p6.font.bold = True
    p6.font.color.rgb = TEXT_WHITE
    p6_sub = tf_t6.add_paragraph()
    p6_sub.text = "Democratizing explainable social engineering defense with an honest engineering foundation."
    p6_sub.font.name = FONT_MAIN
    p6_sub.font.size = Pt(14)
    p6_sub.font.color.rgb = TEXT_MUTED

    # 3-Column Layout
    col_w = Inches(3.75)
    gap_col = Inches(0.24)
    c_start = Inches(0.8)

    # Left: Impact
    c1_x = c_start
    add_card(s6, c1_x, Inches(1.85), col_w, Inches(3.8), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)
    tb_c1 = s6.shapes.add_textbox(c1_x + Inches(0.2), Inches(2.0), col_w - Inches(0.4), Inches(3.5))
    tfc1 = tb_c1.text_frame
    tfc1.word_wrap = True
    tfc1.margin_left = tfc1.margin_top = tfc1.margin_right = tfc1.margin_bottom = 0

    pc1_h = tfc1.paragraphs[0]
    pc1_h.text = "USER IMPACT"
    pc1_h.font.name = FONT_MONO
    pc1_h.font.size = Pt(11)
    pc1_h.font.bold = True
    pc1_h.font.color.rgb = SKY

    impacts = [
        ("Pre-Click Verification", "Helps everyday users spot deceptive attacks before interacting or clicking."),
        ("Transparent Reasoning", "Replaces opaque blocks with clear, educational explanations of why a message is risky."),
        ("Actionable Guidance", "Prescribes exact safe countermeasures to protect non-technical families and seniors."),
    ]
    for h, b in impacts:
        p_i = tfc1.add_paragraph()
        p_i.text = f"• {h}\n"
        p_i.font.name = FONT_MAIN
        p_i.font.size = Pt(11.5)
        p_i.font.bold = True
        p_i.font.color.rgb = TEXT_WHITE
        p_i.space_before = Pt(10)

        run = p_i.add_run()
        run.text = b
        run.font.size = Pt(10.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_SILVER

    # Center: Current Limitations (Engineering Honesty)
    c2_x = c_start + col_w + gap_col
    add_card(s6, c2_x, Inches(1.85), col_w, Inches(3.8), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)
    tb_c2 = s6.shapes.add_textbox(c2_x + Inches(0.2), Inches(2.0), col_w - Inches(0.4), Inches(3.5))
    tfc2 = tb_c2.text_frame
    tfc2.word_wrap = True
    tfc2.margin_left = tfc2.margin_top = tfc2.margin_right = tfc2.margin_bottom = 0

    pc2_h = tfc2.paragraphs[0]
    pc2_h.text = "CURRENT LIMITATIONS"
    pc2_h.font.name = FONT_MONO
    pc2_h.font.size = Pt(11)
    pc2_h.font.bold = True
    pc2_h.font.color.rgb = AMBER

    limits = [
        ("No Active Web Crawling", "Does not visit live web hosts or run browser sandboxing (intentional safety design)."),
        ("OCR Resolution Dependent", "Transcription fidelity relies on image clarity, lighting, and resolution."),
        ("Connectivity Dependent", "Live Gemini reasoning requires internet; offline fallback provides rule-based coverage."),
    ]
    for h, b in limits:
        p_l = tfc2.add_paragraph()
        p_l.text = f"• {h}\n"
        p_l.font.name = FONT_MAIN
        p_l.font.size = Pt(11.5)
        p_l.font.bold = True
        p_l.font.color.rgb = TEXT_WHITE
        p_l.space_before = Pt(10)

        run = p_l.add_run()
        run.text = b
        run.font.size = Pt(10.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_SILVER

    # Right: Future Roadmap
    c3_x = c_start + 2 * (col_w + gap_col)
    add_card(s6, c3_x, Inches(1.85), col_w, Inches(3.8), bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1)
    tb_c3 = s6.shapes.add_textbox(c3_x + Inches(0.2), Inches(2.0), col_w - Inches(0.4), Inches(3.5))
    tfc3 = tb_c3.text_frame
    tfc3.word_wrap = True
    tfc3.margin_left = tfc3.margin_top = tfc3.margin_right = tfc3.margin_bottom = 0

    pc3_h = tfc3.paragraphs[0]
    pc3_h.text = "FUTURE ROADMAP"
    pc3_h.font.name = FONT_MONO
    pc3_h.font.size = Pt(11)
    pc3_h.font.bold = True
    pc3_h.font.color.rgb = CYAN

    roadmap = [
        ("Phase 1: Coverage", "Multilingual scam detection templates and enhanced OCR preprocessing for degraded images."),
        ("Phase 2: Intelligence", "Expanded domain lexical telemetry and on-device small language models for full offline AI."),
        ("Phase 3: Integration", "Security-team dashboard, enterprise API endpoints, and browser extension for real-time triage."),
    ]
    for h, b in roadmap:
        p_r = tfc3.add_paragraph()
        p_r.text = f"• {h}\n"
        p_r.font.name = FONT_MAIN
        p_r.font.size = Pt(11.5)
        p_r.font.bold = True
        p_r.font.color.rgb = TEXT_WHITE
        p_r.space_before = Pt(10)

        run = p_r.add_run()
        run.text = b
        run.font.size = Pt(10.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_SILVER

    # Bottom Closing Statement Card
    add_card(s6, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.0), bg_color=CARD_BG_ALT, border_color=CYAN, border_width=1.5)
    tb_close = s6.shapes.add_textbox(Inches(1.1), Inches(5.88), Inches(11.133), Inches(0.85))
    tf_cl = tb_close.text_frame
    tf_cl.word_wrap = True
    tf_cl.margin_left = tf_cl.margin_top = tf_cl.margin_right = tf_cl.margin_bottom = 0

    pcl1 = tf_cl.paragraphs[0]
    pcl1.text = "“ScamShield AI doesn't just flag a message."
    pcl1.font.name = FONT_MAIN
    pcl1.font.size = Pt(16)
    pcl1.font.bold = True
    pcl1.font.color.rgb = TEXT_WHITE

    pcl2 = tf_cl.add_paragraph()
    pcl2.text = "It helps the user understand the risk and decide what to do next.”"
    pcl2.font.name = FONT_MAIN
    pcl2.font.size = Pt(16)
    pcl2.font.bold = True
    pcl2.font.color.rgb = CYAN

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")


if __name__ == "__main__":
    build_presentation()
